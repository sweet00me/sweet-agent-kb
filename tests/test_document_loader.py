
"""
文件：tests/test_document_loader.py

用途：
    测试统一文档解析器。

不会写入真实知识库。
不会调用 Embedding。
不会修改 ChromaDB。
"""

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from app.rag.document_loader import (
    load_document,
    DocumentLoadError,
)


def main():

    with TemporaryDirectory() as temp:

        root = Path(temp) / "documents"
        notes = root / "notes"
        notes.mkdir(parents=True)

        # =============================
        # 1. Markdown
        # =============================

        (notes / "research.md").write_text(
            "# SD-LoRA 实验\n"
            "ADE20K mIoU: 50.37%",
            encoding="utf-8",
        )

        result = load_document(
            "notes/research.md",
            root,
        )

        assert result["file_type"] == "md"
        assert "50.37%" in result["content"]

        first_id = result["document_id"]

        again = load_document(
            "notes/research.md",
            root,
        )

        assert first_id == again["document_id"]

        print("PASS：Markdown 解析与稳定 ID")

        # =============================
        # 2. TXT
        # =============================

        (notes / "note.txt").write_text(
            "这是中文科研笔记。",
            encoding="utf-8",
        )

        result = load_document(
            "notes/note.txt",
            root,
        )

        assert result["file_type"] == "txt"
        assert "科研笔记" in result["content"]

        print("PASS：TXT 中文解析")

        # =============================
        # 3. Word
        # =============================

        doc = Document()

        doc.add_heading("实验结果", level=1)

        doc.add_paragraph(
            "SD-LoRA 在 ADE20K 上进行实验。"
        )

        table = doc.add_table(
            rows=2,
            cols=2,
        )

        table.cell(0, 0).text = "方法"
        table.cell(0, 1).text = "mIoU"
        table.cell(1, 0).text = "SD-LoRA"
        table.cell(1, 1).text = "50.37"

        doc.save(notes / "experiment.docx")

        result = load_document(
            "notes/experiment.docx",
            root,
        )

        assert result["file_type"] == "docx"
        assert "实验结果" in result["content"]
        assert "SD-LoRA" in result["content"]
        assert "50.37" in result["content"]

        print("PASS：Word 正文及表格解析")

        # =============================
        # 4. 安全检查
        # =============================

        try:

            load_document(
                "../secret.md",
                root,
            )

            raise AssertionError(
                "路径穿越未被阻止"
            )

        except DocumentLoadError:
            pass

        print("PASS：路径穿越防护")

    print("\n统一文档解析测试全部通过")


if __name__ == "__main__":
    main()
