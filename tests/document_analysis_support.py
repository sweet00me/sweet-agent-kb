"""实际解析器和 SQLite 的临时文档夹具；禁止测试访问网络。"""
from datetime import datetime
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from app.rag.document_sync import sync_storage


NOTES = "KnowledgeBase/documents/notes"
DIARY = "KnowledgeBase/documents/Diary/24.7.16.md"


class AnalysisFixture(unittest.TestCase):
    def setUp(self):
        self.enterContext(patch("socket.socket.connect", side_effect=AssertionError("offline test attempted network")))
        self.root = Path(self.enterContext(tempfile.TemporaryDirectory())) / "Storage"
        self.root.mkdir()
        self.parsed = self.root / "KnowledgeBase/parsed_documents"
        self.db = self.root / "KnowledgeBase/sqlite/document_sync.sqlite3"
        self.snapshots = self.root / "KnowledgeBase/sqlite/document_catalog.sqlite3"
        self.write(NOTES + "/alpha.md", "# Alpha\n完整正文，不截取摘要。\n", "2024-01-01T17:00:00+00:00")
        self.write(NOTES + "/nested/beta.txt", "中文 English 😀\n第二段。\n", "2024-01-03T01:00:00+00:00")
        self.write("KnowledgeBase/documents/notes-other/other.md", "兄弟目录不属于 notes。", "2024-01-02T01:00:00+00:00")
        self.write(DIARY, "EXCLUDED_DIARY_BODY", "2024-01-02T01:00:00+00:00")
        self.write("private/secret.md", "PRIVATE_BODY", "2024-01-03T01:00:00+00:00")
        self.write("PC_Files/Uploads/readme.txt", "上传文档正文。", "2024-01-04T01:00:00+00:00")
        path = self.root / (NOTES + "/report.docx")
        doc = Document()
        doc.add_paragraph("Word 正文第一段。")
        doc.add_paragraph("Word 正文最后一段。")
        doc.save(path)
        self.set_date(path, "2024-01-04T01:00:00+00:00")
        self.sync()

    def set_date(self, path, value):
        seconds = datetime.fromisoformat(value).timestamp()
        os.utime(path, (seconds, seconds))

    def write(self, relative, text, modified="2024-01-02T01:00:00+00:00"):
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        self.set_date(path, modified)
        return path

    def sync(self):
        result = sync_storage(self.root, self.parsed, self.db, commit=True)
        self.assertEqual(result["failed"], 0)

    def catalog(self, **kwargs):
        from app.document_analysis.catalog import DocumentCatalog
        return DocumentCatalog(self.root, **kwargs)

    def cached(self, source):
        with sqlite3.connect(self.db) as connection:
            path = connection.execute("SELECT parsed_path FROM parsed_documents WHERE source_path=?", (source,)).fetchone()[0]
        target = Path(path)
        return target, json.loads(target.read_text())
